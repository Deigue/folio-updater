"""Tests for `folio generate`: the whole folio, written as one workbook."""

from __future__ import annotations

import os
from typing import TYPE_CHECKING, NoReturn

from openpyxl import load_workbook

from cli.main import app
from db import drop_table, get_columns, get_connection, get_row_count
from db.backup import backup_file
from domain import Table
from engine.ledger import LedgerUnavailableError
from exporters.excel_writer import sheet_name
from importers.excel_importer import import_transactions

from .helpers.cli import assert_cli_success, assert_in_output, run_cli_with_config
from .helpers.seed import seed_fx, seed_transaction

if TYPE_CHECKING:
    from pathlib import Path

    import pytest

    from app.app_context import AppContext

    from .helpers.cli import CliTestResult
    from .test_types import TempContext

FX = {
    "2025-08-14": "1.2500",
    "2025-08-15": "1.2500",
    "2025-08-18": "1.2500",
    "2025-08-19": "1.2500",
    "2025-08-20": "1.2500",
}

# Every sheet a folio holding in two account types, in two accounts, gets.
EVERY_SHEET = [
    "Summary",
    "Flows",
    "Performance",
    "Portfolio",
    "NON-REGISTERED",
    "TFSA",
    "WS-PERSONAL",
    "WS-TFSA",
    "Ledger",
    "Cost Base",
    "Txns",
    "FX",
    "Tickers",
]


def _seed_folio() -> None:
    """One USD holding in a TFSA, another in a non-registered account."""
    seed_fx(FX)
    seed_transaction(
        action="CONTRIBUTION",
        account="WS-TFSA",
        ticker=None,
        currency="CAD",
        amount="10000",
        price=None,
        units=None,
        date="2025-08-14",
    )
    seed_transaction(
        account="WS-TFSA",
        ticker="TESTTKR",
        amount="-1000",
        price="100",
        units="10",
    )
    seed_transaction(
        account="WS-PERSONAL",
        ticker="OTHER",
        amount="-400",
        price="40",
        units="10",
    )


def _generate(ctx: AppContext, target: Path, *extra: str) -> CliTestResult:
    """Run `folio generate` into `target`, with any extra arguments."""
    return run_cli_with_config(
        ctx.config,
        app,
        ["generate", "--out", str(target), *extra],
    )


def test_generate_writes_a_sheet_for_every_view_of_the_folio(
    temp_ctx: TempContext,
    tmp_path: Path,
) -> None:
    """Only pools that hold something get a dashboard: no RRSP sheet here."""
    with temp_ctx() as ctx:
        _seed_folio()
        target = tmp_path / "folio.xlsx"

        result = _generate(ctx, target)

        assert_cli_success(result)
        assert load_workbook(target).sheetnames == EVERY_SHEET


def test_the_txns_sheet_imports_back_into_an_empty_folio(
    temp_ctx: TempContext,
    tmp_path: Path,
) -> None:
    """The workbook is a backup a folio can be rebuilt from.

    The import pipeline grafts any column it does not recognise onto the
    database, so a Txns sheet carrying a computed column would widen the table
    the first time it came back in. It must carry exactly what was stored.
    """
    with temp_ctx() as ctx:
        _seed_folio()
        target = tmp_path / "folio.xlsx"
        assert_cli_success(_generate(ctx, target))

        with get_connection() as conn:
            stored = set(get_columns(conn, Table.TXNS))
            count = get_row_count(conn, Table.TXNS)
            drop_table(conn, Table.TXNS)

        # Out of the tab strip, since the Ledger shows a reader every column,
        # but still found by name.
        assert load_workbook(target)[ctx.config.txn_sheet].sheet_state == "hidden"
        import_transactions(target, None, ctx.config.txn_sheet)

        with get_connection() as conn:
            assert set(get_columns(conn, Table.TXNS)) == stored
            assert get_row_count(conn, Table.TXNS) == count


def test_only_narrows_the_workbook_to_the_sections_named(
    temp_ctx: TempContext,
    tmp_path: Path,
) -> None:
    with temp_ctx() as ctx:
        _seed_folio()
        target = tmp_path / "folio.xlsx"

        # Spacing and a stray trailing comma are forgiven.
        result = _generate(ctx, target, "--only", "acb, summary,")

        assert_cli_success(result)
        assert load_workbook(target).sheetnames == [
            "Summary",
            "Flows",
            "Ledger",
            "Cost Base",
        ]


def test_only_stored_writes_just_the_stored_sheets(
    temp_ctx: TempContext,
    tmp_path: Path,
) -> None:
    with temp_ctx() as ctx:
        _seed_folio()
        target = tmp_path / "folio.xlsx"

        result = _generate(ctx, target, "--only", "stored")

        assert_cli_success(result)
        assert load_workbook(target).sheetnames == ["Txns", "FX", "Tickers"]


def test_generating_over_the_folio_keeps_a_backup_of_the_last_one(
    temp_ctx: TempContext,
) -> None:
    """The configured folio is rotated into its backups before it is overwritten."""
    backup = {"enabled": True, "path": "backups", "max_backups": 5}
    with temp_ctx(backup=backup) as ctx:
        _seed_folio()
        config = ctx.config

        for _ in range(2):
            assert_cli_success(run_cli_with_config(config, app, ["generate"]))

        kept = list((config.backup_path / "folio_xlsx").glob("*.xlsx"))
        assert len(kept) == 1


def test_rotation_keeps_the_newest_backups_by_name_not_by_modified_time(
    temp_ctx: TempContext,
) -> None:
    """A folio restored by hand has an old modified time, yet its backup is newest."""
    backup = {"enabled": True, "path": "backups", "max_backups": 2}
    with temp_ctx(backup=backup) as ctx:
        config = ctx.config
        config.folio_path.parent.mkdir(parents=True, exist_ok=True)
        config.folio_path.write_bytes(b"restored")
        os.utime(config.folio_path, (0, 0))
        kept_dir = config.backup_path / "folio_xlsx"
        kept_dir.mkdir(parents=True)
        for stamp in ("20200101_000000", "20200102_000000"):
            (kept_dir / f"folio_{stamp}.xlsx").write_bytes(b"earlier")

        backup_file(config.folio_path)

        names = sorted(path.name for path in kept_dir.glob("*.xlsx"))
        assert len(names) == 2
        assert names[0] == "folio_20200102_000000.xlsx"
        assert names[1] > names[0]  # the backup just taken


def test_only_refuses_a_section_that_does_not_exist(
    temp_ctx: TempContext,
    tmp_path: Path,
) -> None:
    with temp_ctx() as ctx:
        _seed_folio()
        target = tmp_path / "folio.xlsx"

        result = _generate(ctx, target, "--only", "dashboards")

        assert result.exit_code == 1
        assert_in_output("summary, perf, dash, accounts, acb, stored", result)
        assert not target.exists()


def test_generate_refuses_a_csv(temp_ctx: TempContext, tmp_path: Path) -> None:
    """A workbook of a dozen sheets has nowhere to go in one table."""
    with temp_ctx() as ctx:
        _seed_folio()
        target = tmp_path / "folio.csv"

        result = _generate(ctx, target)

        assert result.exit_code == 1
        assert_in_output("A CSV holds one table", result)
        assert not target.exists()


def test_sheet_names_differing_only_in_case_are_told_apart() -> None:
    """Excel matches sheet names without case, so the second would be the first."""
    renamed = sheet_name("WS-TFSA", ["WS-tfsa"])

    assert renamed.casefold() != "ws-tfsa"


def test_an_unreplayable_folio_still_writes_its_stored_sheets(
    temp_ctx: TempContext,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`generate` wrote the transactions before it could value anything.

    A folio the engine cannot replay must not lose that: it gets the stored
    sheets, and a warning saying why the rest are missing.
    """

    def unreplayable(**_kwargs: object) -> NoReturn:
        msg = "FX rates are empty, so USD cannot be converted."
        raise LedgerUnavailableError(msg)

    with temp_ctx() as ctx:
        _seed_folio()
        monkeypatch.setattr("cli.commands.generate.open_ledger", unreplayable)
        target = tmp_path / "folio.xlsx"

        result = _generate(ctx, target)

        assert_cli_success(result)
        assert_in_output("only the stored", result)
        assert load_workbook(target).sheetnames == ["Txns", "FX", "Tickers"]


def test_an_empty_folio_writes_nothing(temp_ctx: TempContext, tmp_path: Path) -> None:
    with temp_ctx() as ctx:
        target = tmp_path / "folio.xlsx"

        result = _generate(ctx, target)

        assert_cli_success(result)
        assert_in_output("No transactions", result)
        assert not target.exists()
