"""Render `Table` descriptions into a styled Excel workbook.

Two layout rules are deliberate. The heading row is frozen and filtered, so a
sheet can be sorted and filtered like the grid it is; and the totals sit **below a
blank row**, outside the filtered range, so re-sorting never scrambles them into
the data.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

from exporters import excel_style as style
from exporters.table import (
    DATA_ROLES,
    NUMERIC,
    SIGNED,
    Fmt,
    Role,
    Table,
    as_number,
    is_blank,
    render,
)

if TYPE_CHECKING:
    from collections.abc import Collection, Iterable, Sequence
    from pathlib import Path

    from openpyxl.styles import Font
    from openpyxl.worksheet.worksheet import Worksheet

    from exporters.table import Block, Col, Row

# Excel's own limits on what a sheet may be called.
_NAME_LIMIT = 31
_ILLEGAL = str.maketrans(dict.fromkeys(r"[]:*?/\\", "-"))
_FALLBACK_NAME = "Sheet"


def sheet_name(name: str, taken: Collection[str] = ()) -> str:
    """Make a scope's label legal, and unique within one workbook.

    Args:
        name: The label the report would like to use.
        taken: Names already given to other sheets.

    Returns:
        A name Excel accepts, never empty and never a duplicate.
    """
    cleaned = name.translate(_ILLEGAL).strip("'").strip() or _FALLBACK_NAME
    cleaned = cleaned[:_NAME_LIMIT]
    if cleaned not in taken:
        return cleaned

    for suffix in range(2, 100):
        marker = f"~{suffix}"
        candidate = f"{cleaned[: _NAME_LIMIT - len(marker)]}{marker}"
        if candidate not in taken:
            return candidate
    return cleaned[: _NAME_LIMIT - 4] + "~999"  # pragma: no cover - 98 clashes


def write_workbook(path: Path, tables: Sequence[Table]) -> None:
    """Write every table as its own sheet of one workbook.

    Args:
        path: Where the workbook is written.
        tables: The tables, in the order their sheets should appear.
    """
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)

    taken: list[str] = []
    for table in tables:
        name = sheet_name(table.name, taken)
        taken.append(name)
        write_sheet(workbook.create_sheet(name), table)

    if not taken:  # pragma: no cover - callers always pass at least one table
        workbook.create_sheet(_FALLBACK_NAME)
    workbook.save(path)


def write_sheet(worksheet: Worksheet, table: Table) -> None:
    """Lay one table out on one worksheet.

    Args:
        worksheet: The sheet to write into, assumed empty.
        table: What to write.
    """
    cursor = _write_blocks(worksheet, table.blocks)
    if table.grouped:
        _write_groups(worksheet, table.columns, cursor)
        cursor += 1
    header_row = cursor
    _write_header(worksheet, table.columns, header_row)

    row_number = header_row
    last_data_row = header_row
    filtered = True
    for row in table.rows:
        row_number += 1
        _write_row(worksheet, table.columns, row, row_number)
        if filtered and row.role in DATA_ROLES:
            last_data_row = row_number
        else:
            # The spacer ends the filtered range, and the totals below it stay
            # out of reach of a re-sort.
            filtered = False

    _write_notes(worksheet, table.notes, row_number + 2)
    _finish(worksheet, table, header_row, last_data_row)


def _write_blocks(worksheet: Worksheet, blocks: Iterable[Block]) -> int:
    """Write the labelled panels above the table, returning the next free row."""
    cursor = 1
    for block in blocks:
        title = worksheet.cell(row=cursor, column=1, value=block.title)
        title.font = style.TITLE_FONT
        cursor += 1
        for line in block.lines:
            label = worksheet.cell(row=cursor, column=1, value=line.label)
            label.font = style.LABEL_FONT
            cell = worksheet.cell(row=cursor, column=2, value=_value(line.value))
            cell.number_format = _number_format(line.fmt, line.value)
            if line.note:
                note = worksheet.cell(row=cursor, column=3, value=line.note)
                note.font = style.MUTED_FONT
            cursor += 1
        cursor += 1
    return cursor


def _write_groups(worksheet: Worksheet, columns: Sequence[Col], row: int) -> None:
    """Draw the merged band naming each family of columns."""
    start = 0
    while start < len(columns):
        group = columns[start].group
        end = start
        while end + 1 < len(columns) and columns[end + 1].group == group:
            end += 1
        if group:
            cell = worksheet.cell(row=row, column=start + 1, value=group)
            cell.font = style.GROUP_FONT
            cell.alignment = style.GROUP_ALIGNMENT
            cell.fill = style.ROW_FILLS[Role.SUBTOTAL]
            if end > start:
                worksheet.merge_cells(
                    start_row=row,
                    start_column=start + 1,
                    end_row=row,
                    end_column=end + 1,
                )
        start = end + 1


def _write_header(worksheet: Worksheet, columns: Sequence[Col], row: int) -> None:
    """Write the heading row, which is what the filter and freeze key on."""
    for index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=row, column=index, value=column.header)
        cell.font = style.HEADER_FONT
        cell.fill = style.HEADER_PATTERN
        cell.alignment = style.HEADER_ALIGNMENT


def _write_row(
    worksheet: Worksheet,
    columns: Sequence[Col],
    row: Row,
    number: int,
) -> None:
    """Write one line, styling each cell by its column's role and the row's."""
    fill = style.ROW_FILLS.get(row.role)
    font = style.ROW_FONTS.get(row.role)
    bold = row.role in style.BOLD_ROLES
    cells = enumerate(zip(columns, row.cells, strict=True), start=1)
    for index, (column, raw) in cells:
        cell = worksheet.cell(row=number, column=index, value=_value(raw))
        cell.number_format = _number_format(column.fmt, raw)
        if fill is not None:
            cell.fill = fill
        if row.role is Role.TOTAL:
            cell.border = style.TOTAL_BORDER
        signed = _sign_font(column.fmt, raw, bold=bold)
        if signed is not None:
            cell.font = signed
        elif font is not None:
            cell.font = font


def _write_notes(worksheet: Worksheet, notes: Sequence[str], row: int) -> None:
    """Write the disclosure lines under the table."""
    for offset, note in enumerate(notes):
        cell = worksheet.cell(row=row + offset, column=1, value=note)
        cell.font = style.NOTE_FONT


def _finish(
    worksheet: Worksheet,
    table: Table,
    header_row: int,
    last_data_row: int,
) -> None:
    """Freeze, filter, size and accent the finished sheet."""
    last_column = get_column_letter(len(table.columns))
    worksheet.freeze_panes = f"B{header_row + 1}"
    if last_data_row > header_row:
        worksheet.auto_filter.ref = f"A{header_row}:{last_column}{last_data_row}"

    for index, column in enumerate(table.columns):
        letter = get_column_letter(index + 1)
        worksheet.column_dimensions[letter].width = _width(index, column, table.rows)

    if table.tab_color:
        worksheet.sheet_properties.tabColor = table.tab_color

    _write_bars(worksheet, table, header_row, last_data_row)


def _write_bars(
    worksheet: Worksheet,
    table: Table,
    header_row: int,
    last_data_row: int,
) -> None:
    """Draw an in-cell bar behind the columns that asked for one."""
    if last_data_row <= header_row:
        return
    rule = DataBarRule(
        start_type="num",
        start_value=0,
        end_type="num",
        end_value=style.BAR_FULL_AT / 100,
        color=style.BAR_COLOR,
        showValue=True,
    )
    for index, column in enumerate(table.columns, start=1):
        if not column.bar:
            continue
        letter = get_column_letter(index)
        span = f"{letter}{header_row + 1}:{letter}{last_data_row}"
        worksheet.conditional_formatting.add(span, rule)


def _value(raw: object) -> object:
    """Read one cell value, turning every flavour of blank into a real one.

    A number is written as the number it is, at full precision: the format
    decides what a reader sees, never what the cell holds.
    """
    if is_blank(raw):
        return None
    number = as_number(raw)
    return str(raw) if number is None else number


def _number_format(fmt: Fmt, raw: object) -> str:
    """Pick the format one cell displays at."""
    if fmt is Fmt.UNITS and _whole(raw):
        return style.WHOLE_UNITS
    return style.NUMBER_FORMATS[fmt]


def _whole(raw: object) -> bool:
    """Report whether a share count has no fraction worth showing."""
    number = as_number(raw)
    return number is not None and float(number) == int(number)


def _sign_font(fmt: Fmt, raw: object, *, bold: bool) -> Font | None:
    """Colour a figure whose sign carries meaning."""
    number = as_number(raw) if fmt in SIGNED else None
    return None if number is None else style.sign_font(float(number), bold=bold)


def _width(index: int, column: Col, rows: Sequence[Row]) -> int:
    """Size a column from the widest thing it has to show.

    The heading is measured with room for its filter button, which Excel draws
    inside the cell and which would otherwise sit on top of the text.
    """
    if column.width is not None:
        return column.width
    longest = len(column.header) + style.FILTER_PADDING
    for row in rows:
        if index < len(row.cells):
            longest = max(
                longest,
                len(render(column.fmt, row.cells[index])) + style.WIDTH_PADDING,
            )
    widest = style.NUMBER_MAX_WIDTH if column.fmt in NUMERIC else style.MAX_WIDTH
    return min(max(longest, style.MIN_WIDTH), widest)
