"""How an exported workbook looks: number formats, fills, fonts, accents.

**Number formats never round the stored value.** They decide how many decimals a
reader sees; the cell underneath keeps the full-precision figure the engine
computed, which is what the formula bar shows.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from domain import AccountType
from exporters.table import Fmt, Role

# --- Number formats ----------------------------------------------------------

_MONEY = "#,##0.00"
_MONEY_SIGNED = '#,##0.00;-#,##0.00;"-"'
_PERCENT_SIGNED = '0.00%;-0.00%;"-"'

NUMBER_FORMATS: dict[Fmt, str] = {
    Fmt.TEXT: "@",
    Fmt.DATE: "@",  # dates are stored as YYYY-MM-DD text, and sort as such
    Fmt.ID: "0",
    Fmt.MONEY: _MONEY,
    Fmt.MONEY_SIGNED: _MONEY_SIGNED,
    Fmt.PRICE: "#,##0.00##",
    Fmt.PRICE_SIGNED: '#,##0.00##;-#,##0.00##;"-"',
    Fmt.UNITS: "#,##0.######",
    Fmt.PERCENT: "0.00%",
    Fmt.PERCENT_SIGNED: _PERCENT_SIGNED,
    Fmt.RATE: "0.0000",
}

# A whole share count wants no decimal separator trailing it, which
# `#,##0.######` would leave behind.
WHOLE_UNITS = "#,##0"

# The three sections of an Excel format are positive, negative and zero.
_SECTIONS = 3


def quiet(number_format: str) -> str:
    """Rewrite a format so that a zero draws nothing.

    Args:
        number_format: The format the column would otherwise use.

    Returns:
        The same format with an empty zero section.
    """
    sections = number_format.split(";")
    if len(sections) >= _SECTIONS:
        return ";".join([*sections[:2], '""', *sections[_SECTIONS:]])
    positive = sections[0]
    return f'{positive};-{positive};""'


# --- Palette -----------------------------------------------------------------

GAIN = "1A7F37"
LOSS = "C0392B"
MUTED = "7F7F7F"
INK = "1F1F1F"
PAPER = "FFFFFF"

HEADER_FILL = "1F3864"
SUBTOTAL_FILL = "DCE6F1"
TOTAL_FILL = "BDD7EE"
BAND_FILL = "EDF2FA"
CLOSED_FILL = "F2F2F2"

# Rich names from `ui/vocabulary.py`, in their nearest web-colour hex.
ACCOUNT_TYPE_TABS: dict[AccountType, str] = {
    AccountType.TFSA: "58C27D",
    AccountType.FHSA: "5FD3BC",
    AccountType.RRSP: "6FA8DC",
    AccountType.RRIF: "5B9BD5",
    AccountType.LIRA: "B0C4DE",
    AccountType.RESP: "D7A0D7",
    AccountType.NON_REGISTERED: "F08080",
    AccountType.MARGIN: "FFA07A",
    AccountType.CORPORATE: "D2B48C",
}

FOLIO_TAB = "1F3864"

# --- Cell styles -------------------------------------------------------------

HEADER_FONT = Font(bold=True, color=PAPER)
HEADER_PATTERN = PatternFill("solid", fgColor=HEADER_FILL)
# Headings never wrap
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=False)

GROUP_FONT = Font(bold=True, color=INK)
GROUP_ALIGNMENT = Alignment(horizontal="center")

TITLE_FONT = Font(bold=True, size=12, color=INK)
LABEL_FONT = Font(bold=True, color=INK)
NOTE_FONT = Font(italic=True, color=MUTED)
MUTED_FONT = Font(color=MUTED)
CLOSED_FONT = Font(italic=True, color=MUTED)
SUBTOTAL_FONT = Font(bold=True, color=INK)
TOTAL_FONT = Font(bold=True, size=11, color=INK)

GAIN_FONT = Font(color=GAIN)
LOSS_FONT = Font(color=LOSS)
GAIN_TOTAL_FONT = Font(bold=True, color=GAIN)
LOSS_TOTAL_FONT = Font(bold=True, color=LOSS)

ROW_FILLS: dict[Role, PatternFill | None] = {
    Role.SUBTOTAL: PatternFill("solid", fgColor=SUBTOTAL_FILL),
    Role.TOTAL: PatternFill("solid", fgColor=TOTAL_FILL),
    Role.CLOSED: PatternFill("solid", fgColor=CLOSED_FILL),
}

ROW_FONTS: dict[Role, Font | None] = {
    Role.MUTED: MUTED_FONT,
    Role.CLOSED: CLOSED_FONT,
    Role.SUBTOTAL: SUBTOTAL_FONT,
    Role.TOTAL: TOTAL_FONT,
}

# Bold rows keep their sign colour, so they need their own coloured font.
BOLD_ROLES = frozenset({Role.SUBTOTAL, Role.TOTAL})

TOTAL_BORDER = Border(top=Side(style="thin", color=HEADER_FILL))

# Drawn down the first column of each band, so a sheet scrolled sideways still
# shows where one family of columns ends and the next begins.
BAND_EDGE = Side(style="medium", color=HEADER_FILL)


def edged(*, left: bool, top: bool) -> Border:
    """Build the border one cell needs, from the edges it sits on."""
    return Border(
        left=BAND_EDGE if left else None,
        top=Side(style="thin", color=HEADER_FILL) if top else None,
    )


# In-cell bar drawn behind a weight, full at the same share the terminal's bar
# fills at. Dark enough that the cell's own text, which a dark theme draws
# white, still reads over it.
BAR_COLOR = "2F6FA8"
BAR_FULL_AT = 40  # percent

# --- Widths ------------------------------------------------------------------

MIN_WIDTH = 6
MAX_WIDTH = 42
WIDTH_PADDING = 2

# A figure column is capped tighter than a text one
NUMBER_MAX_WIDTH = 14

# A filtered heading carries a dropdown button at its right edge. Reserving a
# little room keeps the heading readable
FILTER_PADDING = 3


def sign_font(value: float, *, bold: bool = False) -> Font | None:
    """Pick the font that colours a figure by its sign.

    Args:
        value: The number behind the cell.
        bold: Whether the row it sits on is drawn bold.

    Returns:
        The coloured font, or None when the figure is zero and says nothing.
    """
    if value > 0:
        return GAIN_TOTAL_FONT if bold else GAIN_FONT
    if value < 0:
        return LOSS_TOTAL_FONT if bold else LOSS_FONT
    return None
